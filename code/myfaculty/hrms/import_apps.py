"""Εισαγωγή συμβάσεων (Application) από Excel.

Η ροή είναι: parse_workbook -> validate_rows -> save_rows. Ο έλεγχος κάθε γραμμής γίνεται
με την ίδια φόρμα που χρησιμοποιεί ο HR στην καταχώρηση μίας σύμβασης, ώστε να ισχύουν
ακριβώς οι ίδιοι κανόνες.
"""
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
import os

from django.conf import settings
from django.core.files import File
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

from hrms.forms import APPLICATION_LABELS
from hrms.models import Application

MAX_ROWS = 500
MAX_FILE_SIZE = 2 * 1024 * 1024
DATE_FORMATS = ('%d/%m/%Y', '%d.%m.%Y', '%d-%m-%Y', '%Y-%m-%d')
SHEET_NAME = 'Συμβάσεις'

# (επικεφαλίδα στο Excel, πεδίο της φόρμας, πλάτος στήλης)
COLUMNS = [
    ('Τύπος', 'type', 38),
    ('Φορέας', 'organization', 30),
    ('Έναρξη', 'start_date', 14),
    ('Λήξη', 'end_date', 14),
    ('Κατάσταση', 'status', 16),
    ('Τύπος εργασίας', 'work_type', 16),
    ('Αρ. σύμβασης', 'contract_id', 18),
    ('Αρ. βεβαίωσης', 'document2_id', 18),
    ('Σημειώσεις', 'notes', 30),
    ('Σημειώσεις υπηρεσίας', 'notes_decision', 30),
    ('Ποσό', 'payment', 12),
    ('Ανάθεση (από 6)', 'assignment_percentage', 16),
    ('Ώρες', 'no_hours', 10),
    ('Ώρες ανά ημέρα', 'ref_hours', 16),
    ('Ημέρες', 'no_days', 10),
    ('Ποσοστό %', 'day_percentage', 12),
]
HEADERS = [c[0] for c in COLUMNS]
HEADER_TO_FIELD = {c[0]: c[1] for c in COLUMNS}

DATE_FIELDS = ('start_date', 'end_date')
NUMBER_FIELDS = ('payment', 'assignment_percentage', 'no_hours', 'ref_hours', 'no_days', 'day_percentage')

# Στήλες που απαιτούνται επιπλέον ανά τύπο σύμβασης (για τις οδηγίες)
TYPE_EXTRA_COLUMNS = {
    Application.FULLTIME: [],
    Application.WORKCONTRACT: ['Ποσό'],
    Application.TEACHING407_FRACTION: ['Ανάθεση (από 6)'],
    Application.TEACHING_HOURS: ['Ώρες', 'Ώρες ανά ημέρα'],
    Application.TEACHING_DAYS: ['Ημέρες'],
    Application.OTHER: ['Ποσοστό %'],
    Application.EXTERNAL: ['Ποσοστό %'],
}

# Οι ετικέτες είναι lazy μεταφράσεις: αντιστρέφονται την ώρα της χρήσης, όχι στο import.
def _by_label(choices):
    return {str(label): code for code, label in choices}


@dataclass
class RowResult:
    row_no: int
    type_text: str = ''
    type_code: str = ''
    data: dict = field(default_factory=dict)
    errors: list = field(default_factory=list)
    form: object = None

    @property
    def valid(self):
        return not self.errors

    @property
    def type_label(self):
        # Για άγνωστο τύπο δείχνουμε το κείμενο όπως το έγραψε ο χρήστης, ώστε να φαίνεται τι είναι λάθος
        return str(dict(Application.TYPE_CHOICES).get(self.type_code, self.type_text))


WIDTHS = {header: width for header, _, width in COLUMNS}


def application_to_row(app):
    """Η σύμβαση ως γραμμή του ίδιου σχήματος με το template εισαγωγής: {επικεφαλίδα: τιμή}."""
    values = {
        'type': app.get_type_display(),
        'organization': app.organization,
        'start_date': app.start_date,
        'end_date': app.end_date,
        'status': app.get_status_display(),
        'work_type': app.get_work_type_display() if app.work_type else '',
        'contract_id': app.contract_id,
        'document2_id': app.document2_id,
        'notes': app.notes,
        'notes_decision': app.notes_decision,
        'payment': app.payment,
        'assignment_percentage': app.assignment_percentage,
        'no_hours': app.no_hours,
        'ref_hours': app.ref_hours,
        'no_days': app.no_days,
        # Το ποσοστό δίνεται από τον χρήστη μόνο στους τύπους OTHER/EXTERNAL (στους υπόλοιπους υπολογίζεται)
        'day_percentage': (round(app.day_percentage * 100, 4)
                           if app.type in (Application.OTHER, Application.EXTERNAL) and app.day_percentage is not None
                           else None),
    }
    return {header: values[name] for header, name, _ in COLUMNS}


LISTS_SHEET = 'Λίστες'


def add_dropdowns(wb, ws):
    """Dropdown στις στήλες με περιορισμένες τιμές. Οι τιμές μπαίνουν σε κρυφό φύλλο, γιατί μια
    inline λίστα του Excel δεν μπορεί να ξεπερνά τους 255 χαρακτήρες (οι τύποι σύμβασης ξεπερνούν)."""
    lists_ws = wb.create_sheet(LISTS_SHEET)
    sources = {
        'type': Application.TYPE_CHOICES,
        'status': Application.STATUS_CHOICES,
        'work_type': Application.WORK_CHOICES,
    }
    last_row = MAX_ROWS + 1
    for list_col, (name, choices) in enumerate(sources.items(), start=1):
        for row, (_, label) in enumerate(choices, start=1):
            lists_ws.cell(row=row, column=list_col, value=str(label))
        list_letter = lists_ws.cell(row=1, column=list_col).column_letter
        formula = "'%s'!$%s$1:$%s$%d" % (LISTS_SHEET, list_letter, list_letter, len(choices))

        validation = DataValidation(type='list', formula1=formula, allow_blank=True, showDropDown=False)
        validation.error = 'Επιλέξτε μία από τις τιμές της λίστας.'
        validation.errorTitle = 'Μη έγκυρη τιμή'
        validation.showErrorMessage = True
        field_names = [c[1] for c in COLUMNS]
        col_letter = ws.cell(row=1, column=field_names.index(name) + 1).column_letter
        validation.add('%s2:%s%d' % (col_letter, col_letter, last_row))
        ws.add_data_validation(validation)
    lists_ws.sheet_state = 'hidden'


def build_template_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    example = {
        'Τύπος': str(Application.TYPE_CHOICES[0][1]),
        'Φορέας': 'Παράδειγμα ΑΕΕ',
        'Έναρξη': '01/09/2015',
        'Λήξη': '30/06/2016',
        'Κατάσταση': str(Application.STATUS_CHOICES[1][1]),
        'Τύπος εργασίας': str(Application.WORK_CHOICES[0][1]),
    }
    ws.append([example.get(h, '') for h in HEADERS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for idx, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = 'A2'

    add_dropdowns(wb, ws)

    help_ws = wb.create_sheet('Οδηγίες')
    rows = [
        ['Οδηγίες συμπλήρωσης'],
        ['Κάθε γραμμή του φύλλου «%s» είναι μία σύμβαση. Η πρώτη γραμμή (επικεφαλίδες) δεν πρέπει να αλλάξει.' % SHEET_NAME],
        ['Οι ημερομηνίες γράφονται ως ΗΗ/ΜΜ/ΕΕΕΕ. Η «Κατάσταση» αν μείνει κενή θεωρείται «%s».' % str(dict(Application.STATUS_CHOICES)[Application.APPROVED])],
        ['Οι εγγραφές γράφονται στη βάση μόνο αν ΟΛΕΣ οι γραμμές είναι σωστές.'],
        [],
        ['Τύπος', 'Επιπλέον υποχρεωτικές στήλες'],
    ]
    for code, label in Application.TYPE_CHOICES:
        rows.append([str(label), ', '.join(TYPE_EXTRA_COLUMNS[code]) or '-'])
    rows += [[], ['Επιτρεπτές τιμές «Κατάσταση»'] + [str(l) for _, l in Application.STATUS_CHOICES],
             ['Επιτρεπτές τιμές «Τύπος εργασίας»'] + [str(l) for _, l in Application.WORK_CHOICES],
             [], ['Το «Ποσοστό %» δίνεται από 0 έως 100 (π.χ. 70 για 70%).']]
    for row in rows:
        help_ws.append(row)
    help_ws['A1'].font = Font(bold=True)
    help_ws.column_dimensions['A'].width = 60
    help_ws.column_dimensions['B'].width = 40

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _cell_to_str(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.strftime('%d/%m/%Y')
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# Εκκρεμές αρχείο εισαγωγής: ένα ανά (χρήστης, εργαζόμενος). Ο φάκελος hrms_imports
# δεν υπάρχει στο can_download() (myprofile/checks.py), άρα δεν κατεβαίνει μέσω web.
def upload_path(user, employee):
    return os.path.join(settings.MEDIA_ROOT, 'hrms_imports', user.username, '%d.xlsx' % employee.pk)

def store_upload(uploaded_file, user, employee):
    path = upload_path(user, employee)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'wb') as f:
        for chunk in uploaded_file.chunks():
            f.write(chunk)

def open_upload(user, employee):
    """Επιστρέφει File του εκκρεμούς αρχείου ή None."""
    path = upload_path(user, employee)
    if not os.path.isfile(path):
        return None
    return File(open(path, 'rb'), name=path)

def delete_upload(user, employee):
    try:
        os.remove(upload_path(user, employee))
    except FileNotFoundError:
        pass


def parse_workbook(uploaded_file):
    """Επιστρέφει (raw_rows, file_errors). Κάθε raw_row: {'row_no', 'values': {επικεφαλίδα: κείμενο}}."""
    if uploaded_file.size > MAX_FILE_SIZE:
        return [], ['Το αρχείο είναι μεγαλύτερο από 2MB.']
    if not uploaded_file.name.lower().endswith('.xlsx'):
        return [], ['Το αρχείο πρέπει να είναι τύπου .xlsx.']

    try:
        wb = load_workbook(uploaded_file, read_only=True, data_only=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        return [], ['Δεν ήταν δυνατή η ανάγνωση του αρχείου. Χρησιμοποιήστε το template.']

    if not rows:
        return [], ['Το αρχείο είναι κενό.']

    header = [_cell_to_str(h) for h in rows[0]]
    missing = [h for h in HEADERS if h not in header]
    if missing:
        return [], ['Λείπουν στήλες από το αρχείο: %s. Χρησιμοποιήστε το template.' % ', '.join(missing)]

    raw_rows = []
    for offset, row in enumerate(rows[1:], start=2):
        values = {}
        for h, cell in zip(header, row):
            if h in HEADER_TO_FIELD:
                values[h] = cell
        if all(_cell_to_str(v) == '' for v in values.values()):
            continue
        raw_rows.append({'row_no': offset, 'values': {h: _cell_to_str(v) for h, v in values.items()}})

    if not raw_rows:
        return [], ['Το αρχείο δεν περιέχει γραμμές με συμβάσεις.']
    if len(raw_rows) > MAX_ROWS:
        return [], ['Το αρχείο έχει περισσότερες από %d γραμμές.' % MAX_ROWS]
    return raw_rows, []


def _normalize_date(text):
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).strftime('%d/%m/%Y')
        except ValueError:
            continue
    return None


def validate_rows(raw_rows, employee, form_map):
    """Ελέγχει τις γραμμές με τη φόρμα του τύπου τους. form_map: {κωδικός τύπου: κλάση φόρμας}."""
    results = []
    for raw in raw_rows:
        result = RowResult(row_no=raw['row_no'])
        values = raw['values']

        type_text = values.get('Τύπος', '')
        result.type_text = type_text
        # Τα δεδομένα της γραμμής φαίνονται στον πίνακα αποτελεσμάτων ακόμα κι αν η γραμμή έχει σφάλμα
        result.data = {HEADER_TO_FIELD[h]: v for h, v in values.items() if HEADER_TO_FIELD[h] != 'type'}
        result.type_code = _by_label(Application.TYPE_CHOICES).get(type_text, '')
        if not result.type_code or result.type_code not in form_map:
            result.errors.append('Άγνωστος τύπος σύμβασης: «%s».' % type_text)
            results.append(result)
            continue

        data = {HEADER_TO_FIELD[h]: v for h, v in values.items() if HEADER_TO_FIELD[h] != 'type'}

        status = data.get('status', '')
        data['status'] = _by_label(Application.STATUS_CHOICES).get(status, status) if status else Application.APPROVED
        if data.get('work_type'):
            data['work_type'] = _by_label(Application.WORK_CHOICES).get(data['work_type'], data['work_type'])

        for name in DATE_FIELDS:
            label = str(APPLICATION_LABELS.get(name, name))
            if not data.get(name):
                result.errors.append('%s: το πεδίο είναι υποχρεωτικό.' % label)
                continue
            normalized = _normalize_date(data[name])
            if not normalized:
                result.errors.append('%s: μη έγκυρη ημερομηνία «%s» (ΗΗ/ΜΜ/ΕΕΕΕ).' % (label, data[name]))
            else:
                data[name] = normalized
        for name in NUMBER_FIELDS:
            if data.get(name):
                data[name] = data[name].replace(',', '.')

        # Το ποσοστό δίνεται 0-100: η φόρμα EXTERNAL διαιρεί η ίδια με 100, η OTHER περιμένει δεκαδικό.
        if result.type_code == Application.OTHER and data.get('day_percentage'):
            try:
                data['day_percentage'] = str(float(data['day_percentage']) / 100)
            except ValueError:
                pass

        result.data = data
        if not result.errors:
            # Η ApplicationForm.clean δεν αντέχει κενές ημερομηνίες, γι' αυτό ελέγχουμε πρώτα τις ημερομηνίες.
            form = form_map[result.type_code](data, employee=employee)
            if form.is_valid():
                result.form = form
            else:
                for name, messages in form.errors.items():
                    label = str(APPLICATION_LABELS.get(name, name)) if name != '__all__' else 'Γενικό'
                    for message in messages:
                        result.errors.append('%s: %s' % (label, message))
        results.append(result)
    return results


def save_rows(results, employee):
    """Γράφει όλες τις γραμμές ή καμία. Απαιτεί αποτελέσματα του validate_rows χωρίς σφάλματα."""
    if not results or any(not r.valid for r in results):
        raise ValueError('Υπάρχουν γραμμές με σφάλματα.')
    with transaction.atomic():
        for r in results:
            instance = r.form.save(commit=False)
            instance.type = r.type_code
            instance.employee = employee
            instance.save()
            r.form.save_m2m()
    return len(results)
