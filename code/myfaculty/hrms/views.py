# views.py
from django.contrib.auth.decorators import login_required, user_passes_test
from django.http import HttpResponse, FileResponse, Http404
import csv
from datetime import date
from calendar import monthrange
from django.shortcuts import render, get_object_or_404, redirect
from hrms.models import Decision, Application, WageClass, RankClass, Reference, PayrollDecision
from myprofile.models import Employee, PersonalInfo
from core.views import HrCreateView, HrDeleteView, HrMultipleListView, HrUpdateView
from .utils import decision_apps_to_context
from hrms import forms
from django.utils.translation import gettext_lazy as _
from django.urls import reverse_lazy
from myprofile.forms import HrPIForm
from myprofile.checks import is_hr_user, is_handled_by_hr
import base64
from django.core.files.base import ContentFile
import uuid
from core.views import Table

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from export.utils import xlsx_response
from hrms import import_apps
from hrms.recognition_report import render_report
from django.contrib import messages
from django.http import HttpResponseBadRequest
from django.utils.http import url_has_allowed_host_and_scheme

PX_PER_DAY = 4
LABEL_WIDTH = 220
URL_APP_FLAG = {
    Application.OTHER : 'other',
    Application.EXTERNAL : 'external',
    Application.FULLTIME : 'fulltime',
    Application.WORKCONTRACT : 'workcontract',
    Application.TEACHING407_FRACTION : '407',
    Application.TEACHING_HOURS : 'hours',
    Application.TEACHING_DAYS : 'days',  
    
}

FORM_APP_FLAG = {
    Application.OTHER : forms.HrOtherAppForm,
    Application.EXTERNAL : forms.HrExternalAppForm,
    Application.FULLTIME : forms.HrFullTimeAppForm,
    Application.WORKCONTRACT : forms.HrWorkContractAppForm,
    Application.TEACHING407_FRACTION : forms.HrTeachingApp407Form,
    Application.TEACHING_HOURS : forms.HrTeachingAppHourlyForm,
    Application.TEACHING_DAYS : forms.HrTeachingAppDaysForm
}

EMP_FORM_APP_FLAG = {
    Application.OTHER : forms.EmpOtherAppForm,
    Application.EXTERNAL : forms.EmpExternalAppForm,
    Application.FULLTIME : forms.EmpFullTimeAppForm,
    Application.WORKCONTRACT : forms.EmpWorkContractAppForm,
    Application.TEACHING407_FRACTION : forms.EmpTeachingApp407Form,
    Application.TEACHING_HOURS : forms.EmpTeachingAppHourlyForm,
    Application.TEACHING_DAYS : forms.EmpTeachingAppDaysForm
}

HEADLINE_FLAG = {
    Application.OTHER : _('Σύμβαση Γενικής Φύσης'),
    Application.EXTERNAL : _('Εμπειρία Εξωτερικού'),
    Application.FULLTIME : _('Σύμβαση Πλήρους Απασχόλησης'),
    Application.WORKCONTRACT : _('Σύμβαση Έργου'),
    Application.TEACHING407_FRACTION : _('Σύμβαση Διδασκαλίας σύμφωνα με το 407/80'),
    Application.TEACHING_HOURS : _('Καταχώρηση Σύμβασης Διδασκαλίας με ώρες'),
    Application.TEACHING_DAYS : _('Καταχώρηση Σύμβασης Διδασκαλίας με ημέρες'),
}


FLAG_TO_TYPE = { v : k for k, v in URL_APP_FLAG.items() }

def process_cropped_image(request, form):
    cropped_data = request.POST.get('cropped_image')

    if cropped_data:
        try:
            format, imgstr = cropped_data.split(';base64,')
            ext = format.split('/')[-1]
            unique_name = f"pic{uuid.uuid4().hex}.{ext}"
            file = ContentFile(
                base64.b64decode(imgstr),
                name=unique_name,
            )    
            form.instance.pic = file
        except Exception as e:
            print('Crop error: ', e)


def safe_next(request, default):
    """Το ?next= της σελίδας αν είναι ασφαλές (ίδιο host), αλλιώς το default."""
    next_url = request.GET.get('next')
    if next_url and url_has_allowed_host_and_scheme(next_url, allowed_hosts={request.get_host()}):
        return next_url
    return default

def table_context(tables):
    context = {
        'tables' : []
    }
    for name, table in enumerate(tables):
        table_context = table.to_context()
        context['tables'].append(table_context)
    return context


@login_required
@user_passes_test(is_hr_user)  
def hr_select_application_type(request):
    if request.method == 'POST':
        form = forms.HrSelectAppplicationType(request.POST, request.FILES)
        if form.is_valid():
            instance = form.save(commit=False)
            employee = instance.employee
            app_type = instance.type
            app_flag = URL_APP_FLAG[app_type]
            return redirect('hrms:hr_create_app_employee', employee_id = employee.id, app_flag = app_flag)
    else:
        form = forms.HrSelectAppplicationType()
    
    return render(request,'hrms/show_object.html', 
                  context={'form' : form, 
                           'headline' : _('Επιλέξτε τύπο σύμβασης και εργαζόμενο'),
                           'back_url' : reverse_lazy('hrms:hr_dashboard')})

@login_required
@user_passes_test(is_hr_user)  
def hr_select_app_emp(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    if request.method == 'POST':
        form = forms.HrSelectAppplicationType(request.POST, request.FILES, employee=employee)
        if form.is_valid():
            instance = form.save(commit=False)
            app_type = instance.type
            app_flag = URL_APP_FLAG[app_type]
            return redirect('hrms:hr_create_app_employee', employee_id = employee.id, app_flag = app_flag)
    else:
        form = forms.HrSelectAppplicationType(employee=employee)
    
    return render(request,'hrms/show_object.html', 
                  context={'form' : form, 
                           'headline' : _('Επιλέξτε τύπο σύμβασης'),
                           'back_url' : reverse_lazy('hrms:hr_dashboard')})



@login_required
@user_passes_test(is_hr_user)  
def hr_create_app(request, employee_id, app_flag):
    app_type = FLAG_TO_TYPE[app_flag]
    form_class = FORM_APP_FLAG[app_type]
    employee = get_object_or_404(Employee, id = employee_id)

    if request.method == 'POST':
        form = form_class(request.POST, request.FILES, employee=employee)
        if form.is_valid():
            instance = form.save(commit=False)
            instance.type = app_type
            instance.employee = employee
            instance.save()
            form.save_m2m()
            return redirect('hrms:hr_employee_overview', pk = employee.pk)
    else:
        form = form_class(employee=employee)

    return render(request, 'hrms/show_object.html',
                  context = { 'headline' : HEADLINE_FLAG[app_type], 'form' : form,
                             'back_url' : safe_next(request, reverse_lazy('hrms:hr_employee_overview', kwargs={'pk': employee.pk})) } )


@login_required
@user_passes_test(is_hr_user)  
def hr_update_app(request, pk):
    application = get_object_or_404(Application, pk = pk)    
    employee = application.employee
    app_type = application.type
    form_class = FORM_APP_FLAG[app_type]
    
    if request.method == 'POST':
        form = form_class(request.POST, request.FILES, employee=employee, instance=application)
        if form.is_valid():
            form.save()
            Decision.recalc_many(Decision.affected_by([application]))
            return redirect('hrms:hr_employee_overview', pk = application.employee.pk)
    else:
        form = form_class(instance = application, employee=employee)

    delete_url = reverse_lazy('hrms:hr_delete_app', kwargs = {'pk' : pk })
    return render(request, 'hrms/show_object.html',
                  context = { 'headline' : HEADLINE_FLAG[app_type], 
                             'form' : form,
                             'back_url' : safe_next(request, reverse_lazy('hrms:hr_employee_overview', kwargs={'pk': employee.pk})),
                             'delete_url' : delete_url,
                             'delete_text' : _('Διαγραφή') } )


@login_required
@user_passes_test(is_hr_user)  
def hr_create_decision(request, pk):
    employee = get_object_or_404(Employee, pk=pk)

    if request.method == 'POST':
        form = forms.HrDecisionForm(request.POST, request.FILES, employee=employee)
        if form.is_valid():
            instance = form.save(commit=False)
            instance.employee = employee
            instance.save()
            form.save_m2m()
            instance.calc()
            instance.save()
            return redirect('hrms:hr_employee_overview', pk = employee.pk)
    else:
        form = forms.HrDecisionForm(employee=employee)

    return render(request, 'hrms/show_object.html',
                  context = { 'headline' : 'Νέα Απόφαση Αναγνώρισης', 'form' : form,
                              'back_url' : safe_next(request, reverse_lazy('hrms:hr_employee_overview', kwargs={'pk': employee.pk})) } )

@login_required
@user_passes_test(is_hr_user)  
def hr_update_decision(request, pk):
    decision = get_object_or_404(Decision, pk=pk)
    if request.method == 'POST':
        form = forms.HrDecisionForm(request.POST, request.FILES, instance=decision)
        if form.is_valid():
            instance = form.save()
            instance.calc()
            instance.save()
            return redirect('hrms:hr_employee_overview', pk = decision.employee.pk)
    else:
        form = forms.HrDecisionForm(instance=decision)
    delete_url = reverse_lazy('hrms:hr_delete_dec', kwargs = {'pk' : decision.pk})
    delete_text = _('Διαγραφή')
    return render(request, 'hrms/show_object.html',
                  context = { 'headline' : 'Επεξεργασία Απόφασης Αναγνώρισης', 
                              'form' : form,
                              'back_url' : safe_next(request, reverse_lazy('hrms:hr_employee_overview', kwargs={'pk': decision.employee.pk})),
                              'delete_url' : delete_url,
                              'delete_text' : delete_text  } )



@login_required
@user_passes_test(is_hr_user)  
def hr_employee_overview(request, pk):

    employee = get_object_or_404(Employee, pk = pk)
    back_url = safe_next(request, reverse_lazy('hrms:hr_employee_list'))
    applications = list(Application.objects.filter(employee=employee).prefetch_related('decisions'))
    for application in applications:
        # Συμβάσεις που ανήκουν σε απόφαση δεν διαγράφονται μαζικά (μόνο μία-μία από τη φόρμα τους)
        if application.decisions.all():
            application.select_disabled_reason = _('Περιλαμβάνεται σε απόφαση')
    decisions = Decision.objects.filter(employee=employee)

    app_table = Table( 
            fields = ['date_submitted', 'organization', 'get_type_display', 'get_status_display', 'start_date', 'end_date'],
            table_title = _('Καταχωρημένες Συμβάσεις'),
            headers = {
                'date_submitted' : 'Ημερομηνία Υποβολής', 
                'organization' : 'Φορέας',
                'get_status_display' : 'Κατάσταση Σύμβασης',
                'start_date' : 'Έναρξη',
                'end_date' : "Λήξη",
                'get_type_display' : "Τύπος"
            },
            table_id = 'applications',
            next = request.path,
            create_button = True,
            create_text = _('Νέα Σύμβαση'),
            create_url = reverse_lazy('hrms:hr_select_app_emp', kwargs= {'pk' : employee.pk}),
            select_checkboxes = True,
            bulk_delete_url = reverse_lazy('hrms:hr_bulk_delete_apps', kwargs={'pk': employee.pk}),
            bulk_delete_confirm = _('Να διαγραφούν οι επιλεγμένες συμβάσεις;'),
            update_text = _('Λεπτομέρειες'),
            update_button_icon = 'visibility', 
            update_url = 'hrms:hr_update_app',
            objects = applications
        )
    
    dec_table = Table( 
            fields = ['date', 'total_days', 'we_as_str', 'ref_no'],
            table_title = _('Αποφάσεις Αναγνώρισης Προυπηρεσίας'),
            headers = {
                'date' : _('Ημερομηνία'), 
                'total_days' : _('Αριθμός Ημερών'),
                'we_as_str' : _('Χρόνος Π/Υ'),
                'ref_no' : _('Αριθμός Αναφοράς')
            },
            table_id = 'decisions',
            next = request.path,
            create_button = True,
            create_text = _('Νέα Απόφαση'),
            create_url = reverse_lazy('hrms:hr_create_decision', kwargs = {'pk' : employee.pk}),
            extra_buttons = True,
            extra_button_icon = 'function',
            extra_text = _('Υπολογισμός'),
            extra_url = 'hrms:hr_show_decision_info',
            update_text = _('Ενημέρωση'),
            update_button_icon = 'visibility', 
            update_url = 'hrms:hr_update_decision',
            objects = decisions
        )

    payroll_decisions = PayrollDecision.objects.filter(employee=employee)

    payroll_table = Table(
            fields = ['decision_date', 'effective_date', 'get_decision_type_display', 'protocol_number'],
            table_title = _('Υπηρεσιακές Μεταβολές'),
            headers = {
                'decision_date' : _('Ημερομηνία Πράξης'),
                'effective_date' : _('Ημερομηνία Ισχύος'),
                'get_decision_type_display' : _('Τύπος'),
                'protocol_number' : _('Αρ. Πρωτοκόλλου'),
            },
            table_id = 'payroll_decisions',
            next = request.path,
            create_button = True,
            create_text = _('Νέα Μεταβολή'),
            create_url = reverse_lazy('hrms:hr_select_payroll_decision_emp', kwargs={'pk': employee.pk}),
            update_text = _('Ενημέρωση'),
            update_button_icon = 'edit',
            update_url = 'hrms:hr_update_payroll_decision',
            objects = payroll_decisions,
        )

    wageclass = get_object_or_404(WageClass, employee=employee)
    rankclass = get_object_or_404(RankClass, employee=employee)
    newer_decision = Decision.objects.filter(employee=employee).order_by('-date')
    if newer_decision.exists():
        newer_decision = newer_decision.first()
        experience = newer_decision.calc()
    else:
        experience = None

    return render(request, 'hrms/hr_employee.html', context = {
        'app_table' : app_table.to_context(),
        'dec_table' : dec_table.to_context(),
        'payroll_table' : payroll_table.to_context(),
        'employee' : employee,
        'wageclass' : wageclass,
        'rankclass' : rankclass,
        'experience' : experience,
        'next' : request.path,
        'back_url' : back_url,
    })

@login_required
@user_passes_test(is_hr_user)  
def hr_dashboard(request):
    return render(request, 'hrms/hr_dashboard.html',
                  context = {'headline_above' : 'Dashboard'})

@login_required
@user_passes_test(is_hr_user)  
def hr_show_decision_info(request, pk):
    decision = get_object_or_404(Decision, pk = pk)
    prior_decisions = decision.get_priors()
    result = decision.calc()
    
    return render(
        request,
        "hrms/show_decision_info.html",
        context= {
            'decision' : decision,
            'prior_decisions' : prior_decisions,
            'back_url' : reverse_lazy('hrms:hr_employee_overview', kwargs = {'pk' : decision.employee.pk})
        } | result
    )

@login_required
@user_passes_test(is_hr_user)  
def hr_delete_app(request, pk):
    app = get_object_or_404(Application, pk=pk)
    employee = app.employee
    decisions = Decision.affected_by([app])
    app.delete()
    Decision.recalc_many(decisions)
    return redirect('hrms:hr_employee_overview', pk=employee.pk)


@login_required
@user_passes_test(is_hr_user)  
def hr_delete_dec(request, pk):
    dec = get_object_or_404(Decision, pk=pk)
    employee = dec.employee
    dec.delete()
    return redirect('hrms:hr_employee_overview', pk=employee.pk)


@login_required
@user_passes_test(is_hr_user)
def hr_select_payroll_decision_type(request):
    if request.method == 'POST':
        form = forms.HrSelectPayrollDecisionType(request.POST)
        if form.is_valid():
            instance = form.save(commit=False)
            return redirect('hrms:hr_create_payroll_decision',
                             employee_id=instance.employee.id, decision_type=instance.decision_type)
    else:
        form = forms.HrSelectPayrollDecisionType()
    return render(request, 'hrms/show_object.html', context={
        'form': form, 'headline': _('Επιλέξτε τύπο απόφασης και εργαζόμενο'),
        'back_url': reverse_lazy('hrms:hr_dashboard')})


@login_required
@user_passes_test(is_hr_user)
def hr_select_payroll_decision_emp(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    if request.method == 'POST':
        form = forms.HrSelectPayrollDecisionType(request.POST, employee=employee)
        if form.is_valid():
            instance = form.save(commit=False)
            return redirect('hrms:hr_create_payroll_decision',
                             employee_id=employee.id, decision_type=instance.decision_type)
    else:
        form = forms.HrSelectPayrollDecisionType(employee=employee)
    return render(request, 'hrms/show_object.html', context={
        'form': form, 'headline': _('Επιλέξτε τύπο απόφασης'),
        'back_url': reverse_lazy('hrms:hr_employee_overview', kwargs={'pk': employee.pk})})


# decision_type -> dedicated Hr form class. Types not listed here still use
# the generic HrPayrollDecisionForm (full manual field set) - they'll get
# their own dedicated forms in a later round.
PAYROLL_DECISION_FORM_CLASSES = {
    PayrollDecision.PRIOR_SERVICE_RECOGNITION: forms.HrPriorServiceRecognitionSystemForm,
    PayrollDecision.PRIOR_SERVICE_RECOGNITION_LEGACY: forms.HrPriorServiceRecognitionLegacyForm,
    PayrollDecision.WAGE_PROGRESSION: forms.HrWageProgressionForm,
    PayrollDecision.INDIVIDUAL_RANKING: forms.HrIndividualRankingForm,
    PayrollDecision.LEGAL_FRAMEWORK_RECLASSIFICATION: forms.HrLegalFrameworkReclassificationForm,
    PayrollDecision.RANK_PROMOTION: forms.HrRankPromotionForm,
    PayrollDecision.GRADE_CHANGE: forms.HrGradeChangeForm,
    PayrollDecision.APPOINTMENT: forms.HrAppointmentForm,
    PayrollDecision.DUTY_ASSUMPTION: forms.HrDutyAssumptionForm,
    PayrollDecision.MASTERS_RELEVANCE_RECOGNITION: forms.HrMastersRelevanceRecognitionForm,
    PayrollDecision.PHD_RELEVANCE_RECOGNITION: forms.HrPhdRelevanceRecognitionForm,
}


def _hr_payroll_decision_form_class(decision_type):
    return PAYROLL_DECISION_FORM_CLASSES.get(decision_type, forms.HrPayrollDecisionForm)


@login_required
@user_passes_test(is_hr_user)
def hr_create_payroll_decision(request, employee_id, decision_type):
    if decision_type not in dict(PayrollDecision.DECISION_TYPE_CHOICES):
        raise Http404
    employee = get_object_or_404(Employee, id=employee_id)
    form_class = _hr_payroll_decision_form_class(decision_type)

    if request.method == 'POST':
        form = form_class(request.POST, request.FILES,
                           employee=employee, decision_type=decision_type)
        if form.is_valid():
            instance = form.save(commit=False)
            instance.employee = employee
            instance.decision_type = decision_type
            if decision_type == PayrollDecision.PRIOR_SERVICE_RECOGNITION:
                instance.sync_from_linked_decision()
            instance.save()
            return redirect('hrms:hr_employee_overview', pk=employee.pk)
    else:
        form = form_class(employee=employee, decision_type=decision_type)

    headline = dict(PayrollDecision.DECISION_TYPE_CHOICES)[decision_type]
    return render(request, 'hrms/show_object.html', context={
        'headline': headline, 'form': form,
        'back_url': safe_next(request, reverse_lazy('hrms:hr_employee_overview', kwargs={'pk': employee.pk}))})


@login_required
@user_passes_test(is_hr_user)
def hr_update_payroll_decision(request, pk):
    payroll_decision = get_object_or_404(PayrollDecision, pk=pk)
    employee = payroll_decision.employee
    form_class = _hr_payroll_decision_form_class(payroll_decision.decision_type)

    if request.method == 'POST':
        form = form_class(request.POST, request.FILES,
                           employee=employee, instance=payroll_decision)
        if form.is_valid():
            instance = form.save(commit=False)
            if payroll_decision.decision_type == PayrollDecision.PRIOR_SERVICE_RECOGNITION:
                instance.sync_from_linked_decision()
            instance.save()
            if employee:
                return redirect('hrms:hr_employee_overview', pk=employee.pk)
            return redirect('hrms:hr_employee_list')
    else:
        form = form_class(instance=payroll_decision, employee=employee)

    headline = dict(PayrollDecision.DECISION_TYPE_CHOICES)[payroll_decision.decision_type]
    return render(request, 'hrms/show_object.html', context={
        'headline': headline, 'form': form,
        'back_url': safe_next(request, reverse_lazy('hrms:hr_employee_overview', kwargs={'pk': employee.pk})
                              if employee else reverse_lazy('hrms:hr_employee_list')),
        'delete_url': reverse_lazy('hrms:hr_delete_payroll_decision', kwargs={'pk': pk}),
        'delete_text': _('Διαγραφή')})


@login_required
@user_passes_test(is_hr_user)
def hr_delete_payroll_decision(request, pk):
    payroll_decision = get_object_or_404(PayrollDecision, pk=pk)
    employee = payroll_decision.employee
    payroll_decision.delete()
    if employee:
        return redirect('hrms:hr_employee_overview', pk=employee.pk)
    return redirect('hrms:hr_employee_list')


@login_required
@user_passes_test(is_handled_by_hr)
def emp_show_payroll_decision(request, pk):
    employee = Employee.objects.get_from_user(request.user)
    payroll_decision = get_object_or_404(PayrollDecision, pk=pk, employee=employee)
    form = forms.EmpPayrollDecisionForm(instance=payroll_decision)
    headline = dict(PayrollDecision.DECISION_TYPE_CHOICES)[payroll_decision.decision_type]
    return render(request, 'hrms/show_object.html', context={
        'headline': headline, 'form': form,
        'back_url': reverse_lazy('hrms:employee_dashboard')})


@login_required
@user_passes_test(is_hr_user)
def hr_employee_list(request):
    admin_staff = Employee.objects.filter(is_admin_staff=True).order_by(
        'admin_staff__personal_info__surname', 'admin_staff__personal_info__given_name')
    admin_table = Table( 
            fields = ['admin_staff.personal_info.given_name', 
                      'admin_staff.personal_info.surname', 
                      'admin_staff.personal_info.email', 'employee_type'],
            table_title = _('Διοικητικό Προσωπικό'),
            headers = {
                'admin_staff.personal_info.given_name' : 'Όνομα', 
                'admin_staff.personal_info.surname' : 'Επώνυμο', 
                'admin_staff.personal_info.email' : 'Email', 
                'employee_type' : 'Κατηγορία'
            },
            table_id = 'admins',
            table_order = '[[1, "asc"], [0, "asc"]]',
            next = request.path,
            create_button = True,
            create_text = _('Προσθήκη Εργαζόμενου'),
            create_url = reverse_lazy('hrms:hr_create_employee', kwargs = {'etype' : 'adm'}),
            update_text = _('Μεταβολές'),
            update_button_icon = 'timer',
            update_url = 'hrms:hr_employee_overview',
            extra_text = _('Λεπτομέρειες'),
            extra_button_icon = 'visibility', 
            extra_url = 'hrms:hr_update_employee',
            extra_buttons = True,
            objects = admin_staff
        )
    faculty_staff = Employee.objects.filter(is_faculty_staff=True).order_by(
        'faculty_staff__surname', 'faculty_staff__given_name')
    faculty_table = Table( 
            fields = ['given_name', 'surname', 'email', 'employee_type'],
            table_title = _('Διδακτικό Προσωπικό'),
            headers = {
                'given_name' : 'Όνομα', 
                'surname' : 'Επώνυμο', 
                'email' : 'Email', 
                'employee_type' : 'Κατηγορία'
            },
            table_id = 'faculty',
            table_order = '[[1, "asc"], [0, "asc"]]',
            next = request.path,
            create_button = True,
            create_text = _('Προσθήκη Εργαζόμενου'),
            create_url = reverse_lazy('hrms:hr_create_employee', kwargs = {'etype' : 'fac'}),
            update_text = _('Μεταβολές'),
            update_button_icon = 'timer',
            update_url = 'hrms:hr_employee_overview',
            extra_text = _('Λεπτομέρειες'),
            extra_button_icon = 'visibility', 
            extra_url = 'hrms:hr_update_employee',
            extra_buttons = True,
            objects = faculty_staff
        )
    return render(request, 'hrms/hr_list_employees.html', 
                  context={'admin' : admin_table.to_context(),
                           'faculty' : faculty_table.to_context(),
                           'back_url' : reverse_lazy('hrms:hr_dashboard') })

@login_required
@user_passes_test(is_hr_user)
def hr_create_employee(request, etype):
    if etype == 'adm':
        admin_staff= True
    elif etype == 'fac':
        admin_staff = False
    else:
        return HttpResponse(403)
    
    next = request.GET.get('next','')
    if request.method == 'POST':
        form  = forms.HrEmployeeForm(request.POST, request.FILES, admin_staff=admin_staff)

        if form.is_valid():
            form.save()
            return redirect('hrms:hr_employee_list')
        else:
            print('Form errors:')
            print(form.errors)
    else:
        form  = forms.HrEmployeeForm(admin_staff=admin_staff)
    
    return render(request, 'hrms/show_object.html', 
                  context={'form' : form, 'back_url' : next})


@login_required
@user_passes_test(is_hr_user)
def hr_update_employee(request, pk):
    next = request.GET.get('next','')
    employee = get_object_or_404(Employee, pk=pk)
    if request.method == 'POST':
        form  = forms.HrEmployeeForm(request.POST, request.FILES, employee=employee)

        if form.is_valid():
            form.save()
            return redirect('hrms:hr_employee_list')
        else:
            print('Form errors:')
            print(form.errors)
    else:
        form  = forms.HrEmployeeForm(employee=employee)
    
    return render(request, 'hrms/show_object.html', 
                  context={'form' : form, 'back_url' : next,
                           'delete_text' :  'Διαγραφή',
                           'delete_url' : reverse_lazy('hrms:hr_delete_employee', 
                                                       kwargs = {'pk' :employee.pk})})

@login_required
@user_passes_test(is_hr_user)
def hr_delete_employee(request, pk):
    next = reverse_lazy('hrms:hr_employee_list')    
    employee = get_object_or_404(Employee, pk=pk)
    if employee.is_admin_staff:
        staff = employee.admin_staff
    else:
        staff = employee.faculty_staff
    
    pi = staff.personal_info
    pi.delete()
    staff.delete()
    employee.delete()
    return redirect(next)

INTERVAL_LABELS = {
    'start' : 'Από',
    'end' : 'Εώς',
    'calendar_days' : 'Ημερολογιακές Ημέρες',
    'value' : 'Ποσοστό Αναγνώρισης',
    'organization' : 'Φορέας',
    'type' : 'Σχέση Εργασίας',
    'notes' : 'Επισημάνσεις',
    'work_type' : 'Είδος Π/Υ',
    'related_desc' : 'Περιγραφή Συννημένων',     
}

INTERVAL_HEADERS = [d for d in INTERVAL_LABELS.keys()]

CELL_ALIGNMENT = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True
)

@login_required
@user_passes_test(is_hr_user)
def hr_export_decision_ints(request, pk):

    decision = get_object_or_404(Decision, pk=pk)
    result = decision.calc()

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="export.xlsx"'
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Intervals"

    headers = [
        'Από',
        'Έως',
        'Ημερολογιακές Ημέρες',
        'Ημέρες προς Αναγνώριση',        
        'Ποσοστό Αναγνώρισης',
        'Φορέας',
        'Σχέση Εργασίας',
        'Επισημάνσεις',
        'Είδος Π/Υ',
        'Περιγραφή Συνημμένων'
    ]

    ws.append(headers)

    # Επικεφαλίδες
    for cell in ws[1]:
        cell.font = Font(bold=True)
        
    bin_apps_data = {}

    for el in result['entire_map']:

        idxs = el['idx']

        for idx in idxs:
            if idx not in bin_apps_data:

                application = Application.objects.get(pk=idx)

                bin_apps_data[idx] = {
                    'organization': application.organization,
                    'type': application.get_type_display(),
                    'work_type': application.get_work_type_display(),
                    'notes': application.notes_decision,
                    'related_desc': application.doc_descr(),
                }

        start_row = ws.max_row + 1

        first = True

        for idx in idxs:

            a = bin_apps_data[idx]

            if first:
                row = [
                    el['start'].strftime("%d/%m/%Y"),
                    el['end'].strftime("%d/%m/%Y"),
                    el['calendar_days'],
                    el['days'],
                    el['value'],
                ]
                first = False
            else:
                row = ['', '', '', '', '']

            row += [
                a['organization'],
                a['type'],
                a['notes'],
                a['work_type'],
                a['related_desc'],
            ]

            ws.append(row)

        end_row = ws.max_row

        # format percentage
        ws.cell(start_row, 5).number_format = '0.00%'

        # merge μόνο αν υπάρχουν πολλές γραμμές
        if end_row > start_row:

            for col in range(1, 6):
                ws.merge_cells(
                    start_row=start_row,
                    start_column=col,
                    end_row=end_row,
                    end_column=col
                )

    # Freeze header
    ws.freeze_panes = "A2"

    # Autofilter
    ws.auto_filter.ref = ws.dimensions

    # Προαιρετικά πλάτη στηλών
    widths = {
        'A': 15,
        'B': 15,
        'C': 18,
        'D': 18,
        'E': 35,
        'F': 20,
        'G': 40,
        'H': 20,
        'I': 40,
        'J': 40,
        
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    # Στοίχιση όλων των κελιών
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = CELL_ALIGNMENT
    wb.save(response)

    return response

@login_required
@user_passes_test(is_hr_user)
def hr_export_decision_minutes(request, pk):
    decision = get_object_or_404(Decision, pk=pk)
    filename = f'recognition_minutes_{decision.employee.surname()}_{decision.date or decision.pk}.docx'
    return FileResponse(
        render_report(decision),
        as_attachment=True,
        filename=filename,
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )

@login_required
@user_passes_test(is_hr_user)
def hr_decision_info_by_year(request, pk):
    decision = get_object_or_404(Decision, pk=pk)
    days_dict = decision.calc_by_year()
    return render(request, 'hrms/decision_by_year.html', 
                  context = { 'days_dict' : days_dict,
                              'decision' : decision,
                              'back_url' : reverse_lazy('hrms:hr_show_decision_info', kwargs = {'pk' : decision.pk}) })

@login_required
@user_passes_test(is_hr_user)
def hr_decision_in_year(request, pk, year):
    decision = get_object_or_404(Decision, pk=pk)
    if year not in decision.get_years():
        raise Http404('Not Found.')
    
    result = decision.calc_in_year(year)
    applications = result['current_apps']
    interval_map = result['entire_map']
    orgs = result['orgs']
    types = result['types']


    print(interval_map)
    return render(request, 'hrms/decision_in_year.html', 
                  context = { 'interval_map' : interval_map,
                              'applications' : applications,
                              'orgs' : orgs,
                              'types' : types,
                              'back_url' : reverse_lazy('hrms:hr_decision_info_by_year', kwargs = {'pk' : decision.pk}),
                                })

@login_required
@user_passes_test(is_handled_by_hr)
def employee_dashboard(request):
    employee = Employee.objects.get_from_user(request.user)
    applications = Application.objects.filter(employee = employee)
    decisions = Decision.objects.filter(employee = employee)

    app_table = Table( 
            fields = ['date_submitted', 'organization', 'get_type_display', 'get_status_display', 'start_date', 'end_date'],
            table_title = _('Καταχωρημένες Συμβάσεις'),
            headers = {
                'date_submitted' : 'Ημερομηνία Υποβολής', 
                'organization' : 'Φορέας',
                'get_status_display' : 'Κατάσταση Σύμβασης',
                'start_date' : 'Έναρξη',
                'end_date' : "Λήξη",
                'get_type_display' : "Τύπος"
            },
            table_id = 'applications',
            next = request.path,
            create_button = True,
            create_text = _('Νέα Σύμβαση'),
            create_url = reverse_lazy('hrms:emp_create_app'),
            update_text = _('Λεπτομέρειες'),
            update_button_icon = 'visibility', 
            update_url = 'hrms:emp_update_app',
            objects = applications
        )
    
    dec_table = Table( 
            fields = ['date', 'total_days', 'we_as_str', 'ref_no'],
            table_title = _('Αποφάσεις Αναγνώρισης Προυπηρεσίας'),
            headers = {
                'date' : _('Ημερομηνία'), 
                'total_days' : _('Αριθμός Ημερών'),
                'we_as_str' : _('Χρόνος Π/Υ'),
                'ref_no' : _('Αριθμός Αναφοράς')
            },
            table_id = 'decisions',
            next = request.path,
            extra_buttons = True,
            extra_button_icon = 'download',
            extra_text = _('Απόφαση'),
            extra_url = 'hrms:hr_show_decision_info',
            update_buttons = False,
            update_text = _('Ενημέρωση'),
            update_button_icon = 'visibility', 
            update_url = 'hrms:hr_update_decision',
            objects = decisions,
            create_button = False,
        )

    payroll_decisions = PayrollDecision.objects.filter(employee=employee)

    payroll_table = Table(
            fields = ['decision_date', 'effective_date', 'get_decision_type_display', 'protocol_number'],
            table_title = _('Υπηρεσιακές Μεταβολές'),
            headers = {
                'decision_date' : _('Ημερομηνία Πράξης'),
                'effective_date' : _('Ημερομηνία Ισχύος'),
                'get_decision_type_display' : _('Τύπος'),
                'protocol_number' : _('Αρ. Πρωτοκόλλου'),
            },
            table_id = 'payroll_decisions',
            next = request.path,
            create_button = False,
            update_text = _('Λεπτομέρειες'),
            update_button_icon = 'visibility',
            update_url = 'hrms:emp_show_payroll_decision',
            objects = payroll_decisions,
        )

    wageclass = get_object_or_404(WageClass, employee=employee)
    rankclass = get_object_or_404(RankClass, employee=employee)
    newer_decision = Decision.objects.filter(employee=employee).order_by('-date')
    if newer_decision.exists():
        newer_decision = newer_decision.first()
        experience = newer_decision.calc()
    else:
        experience = None

    return render(request, 'hrms/employee.html', context = {
        'app_table' : app_table.to_context(),
        'dec_table' : dec_table.to_context(),
        'payroll_table' : payroll_table.to_context(),
        'employee' : employee,
        'wageclass' : wageclass,
        'rankclass' : rankclass,
        'experience' : experience,
        'next' : request.path,
    })


@login_required
@user_passes_test(is_handled_by_hr)  
def emp_select_application_type(request):
    next = request.GET.get('next', reverse_lazy('hrms:employee_dashboard'))
    if request.method == 'POST':
        form = forms.EmpSelectAppplicationType(request.POST, request.FILES)
        if form.is_valid():
            instance = form.save(commit=False)
            app_type = instance.type
            app_flag = URL_APP_FLAG[app_type]
            return redirect('hrms:emp_create_app_employee', app_flag = app_flag)
    else:
        form = forms.EmpSelectAppplicationType()
    
    return render(request,'hrms/show_object.html', 
                  context={'form' : form, 
                           'headline' : _('Επιλέξτε τύπο σύμβασης και εργαζόμενο'),
                           'back_url' : next})
@login_required
@user_passes_test(is_handled_by_hr)  
def emp_update_app(request, pk):

    next = request.GET.get('next', reverse_lazy('hrms:employee_dashboard'))
    employee = Employee.objects.get_from_user(request.user)
    application = get_object_or_404(Application, pk = pk, employee = employee)    
    app_type = application.type
    form_class = EMP_FORM_APP_FLAG[app_type]
    
    if request.method == 'POST':
        form = form_class(request.POST, request.FILES, employee=employee, instance=application)
        if form.is_valid():
            form.save()
            Decision.recalc_many(Decision.affected_by([application]))
            return redirect(next)
    else:
        form = form_class(instance = application, employee=employee)

    if application.status == Application.PENDING:
        delete_url = reverse_lazy('hrms:emp_delete_app', kwargs = {'pk' : pk })
    else:
        delete_url = None

    return render(request, 'hrms/show_object.html',
                  context = { 'headline' : HEADLINE_FLAG[app_type], 
                             'form' : form,
                             'back_url' : next,
                             'delete_url' : delete_url,
                             'delete_text' : _('Διαγραφή') } )

@login_required
@user_passes_test(is_handled_by_hr)  
def emp_create_app(request, app_flag):
    next = request.GET.get('next', reverse_lazy('hrms:employee_dashboard'))
    app_type = FLAG_TO_TYPE[app_flag]
    form_class = EMP_FORM_APP_FLAG[app_type]
    employee = Employee.objects.get_from_user(request.user)
    
    if request.method == 'POST':
        form = form_class(request.POST, request.FILES, employee=employee)
        if form.is_valid():
            instance = form.save(commit=False)
            instance.type = app_type
            instance.status = Application.PENDING
            instance.employee = employee
            instance.save()
            form.save_m2m()
            return redirect('hrms:employee_dashboard')
    else:
        form = form_class(employee=employee)

    return render(request, 'hrms/show_object.html',
                  context = { 'headline' : HEADLINE_FLAG[app_type], 
                              'form' : form,
                              'back_url' : next } )
   
@login_required
@user_passes_test(is_handled_by_hr)  
def emp_delete_app(request, pk):
    next = request.GET.get('next', reverse_lazy('hrms:employee_dashboard'))
    employee = Employee.objects.get_from_user(request.user)
    app = get_object_or_404(Application, pk=pk, employee = employee, status = Application.PENDING)    
    decisions = Decision.affected_by([app])
    app.delete()
    Decision.recalc_many(decisions)
    return redirect(next)

@login_required
@user_passes_test(is_handled_by_hr)  
def emp_view_decision(request, pk):
    employee = Employee.objects.get_from_user(request.user)
    dec = get_object_or_404(Decision, pk=pk, employee = employee)
    if not dec.document:
        return render(request, 'hrms/message.html',
                      context = {
                        'html_message' : "messages/not_ready_decision.html",
                        'back_url' : reverse_lazy('hrms:employee_dashboard')
                      })
    else:
        return FileResponse(
        dec.my_file.open("rb"),
        as_attachment=True,
        filename=dec.document.name.split("/")[-1],
    )


@login_required
@user_passes_test(is_hr_user)  
def hr_export_applications(request, pk):
    decision = get_object_or_404(Decision, pk=pk)
    objs = Application.objects.filter(decisions = decision).order_by('-start_date')
    # Ίδιες στήλες με το template εισαγωγής συμβάσεων
    data = [import_apps.application_to_row(obj) for obj in objs]

    return xlsx_response(data, headers=import_apps.HEADERS, widths=import_apps.WIDTHS)

@login_required
@user_passes_test(is_hr_user)  
def hr_export_apps_emp(request, pk):
    employee = get_object_or_404(Employee, pk = pk)
    objs = Application.objects.filter(employee = employee).order_by('-start_date')
    # Ίδιες στήλες με το template εισαγωγής συμβάσεων
    data = [import_apps.application_to_row(obj) for obj in objs]

    return xlsx_response(data, headers=import_apps.HEADERS, widths=import_apps.WIDTHS)


##############################################
# Import συμβάσεων από Excel
##############################################
@login_required
@user_passes_test(is_hr_user)
def hr_apps_template(request):
    response = HttpResponse(
        import_apps.build_template_workbook(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="template_symvaseis.xlsx"'
    return response

@login_required
@user_passes_test(is_hr_user)
def hr_upload_apps(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    return render(request, 'hrms/upload_apps.html', context = {
        'employee' : employee,
        'back_url' : reverse_lazy('hrms:hr_employee_overview', kwargs={'pk': employee.pk}),
    })

@login_required
@user_passes_test(is_hr_user)
def hr_validate_apps(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    upload_url = reverse_lazy('hrms:hr_upload_apps', kwargs={'pk': employee.pk})
    if request.method != 'POST':
        return redirect(upload_url)

    uploaded = request.FILES.get('file')
    if not uploaded:
        return render(request, 'hrms/upload_apps.html', context = {
            'employee' : employee, 'back_url' : reverse_lazy('hrms:hr_employee_overview', kwargs={'pk': employee.pk}),
            'file_errors' : ['Δεν επιλέχθηκε αρχείο.'] })

    raw_rows, file_errors = import_apps.parse_workbook(uploaded)
    results = []
    if not file_errors:
        results = import_apps.validate_rows(raw_rows, employee, FORM_APP_FLAG)

    all_valid = bool(results) and all(r.valid for r in results)
    if all_valid:
        import_apps.store_upload(uploaded, request.user, employee)
    else:
        import_apps.delete_upload(request.user, employee)

    return render(request, 'hrms/validate_apps.html', context = {
        'employee' : employee,
        'results' : results,
        'file_errors' : file_errors,
        'all_valid' : all_valid,
        'back_url' : upload_url,
    })

@login_required
@user_passes_test(is_hr_user)
def hr_import_apps(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    if request.method != 'POST':
        return redirect('hrms:hr_upload_apps', pk=employee.pk)

    stored = import_apps.open_upload(request.user, employee)
    if not stored:
        messages.error(request, 'Η εισαγωγή έχει λήξει. Ανεβάστε ξανά το αρχείο.')
        return redirect('hrms:hr_upload_apps', pk=employee.pk)

    # Ξαναελέγχουμε πριν γράψουμε: το αποθηκευμένο αρχείο δεν θεωρείται έγκυρο εκ των προτέρων.
    try:
        raw_rows, file_errors = import_apps.parse_workbook(stored)
    finally:
        stored.close()
        import_apps.delete_upload(request.user, employee)
    if file_errors:
        messages.error(request, 'Η εισαγωγή απέτυχε: ' + ' '.join(file_errors))
        return redirect('hrms:hr_upload_apps', pk=employee.pk)

    results = import_apps.validate_rows(raw_rows, employee, FORM_APP_FLAG)
    try:
        count = import_apps.save_rows(results, employee)
    except ValueError:
        messages.error(request, 'Η εισαγωγή απέτυχε: υπάρχουν γραμμές με σφάλματα. Δεν γράφτηκε καμία σύμβαση.')
        return redirect('hrms:hr_upload_apps', pk=employee.pk)

    messages.success(request, 'Καταχωρήθηκαν %d συμβάσεις.' % count)
    return redirect('hrms:hr_employee_overview', pk=employee.pk)

@login_required
@user_passes_test(is_hr_user)
def hr_bulk_delete_apps(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    if request.method != 'POST':
        return HttpResponseBadRequest()

    ids = [i for i in request.POST.getlist('ids') if i.isdigit()]
    # Μόνο συμβάσεις του συγκεκριμένου εργαζομένου, και όσες δεν ανήκουν σε απόφαση
    requested = Application.objects.filter(pk__in=ids, employee=employee)
    applications = requested.filter(decisions__isnull=True)
    _total, per_model = applications.delete()
    count = per_model.get('hrms.Application', 0)
    skipped = requested.count()
    if count:
        messages.success(request, 'Διαγράφηκαν %d συμβάσεις.' % count)
    elif not skipped:
        messages.warning(request, 'Δεν διαγράφηκε καμία σύμβαση.')
    if skipped:
        messages.warning(request, '%d συμβάσεις δεν διαγράφηκαν επειδή περιλαμβάνονται σε απόφαση.' % skipped)

    next_url = request.POST.get('next')
    if next_url and url_has_allowed_host_and_scheme(next_url, allowed_hosts={request.get_host()}):
        return redirect(next_url)
    return redirect('hrms:hr_employee_overview', pk=employee.pk)
