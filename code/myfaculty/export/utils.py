from docx import Document
from django.conf import settings
import os
from docxtpl import DocxTemplate
from docxcompose.composer import Composer
import io
from pathlib import Path
from django.template.loader import render_to_string
from weasyprint import HTML
from pypdf import PdfWriter, PdfReader
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from django.http import HttpResponse
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from pypdf.generic import NameObject, ArrayObject
import logging

logger = logging.getLogger(__name__)

def convert_to_docx(template, output_file, d):
    doc = DocxTemplate(template)
    doc.render(d)
    doc.save(output_file)

def convert_multiple_to_docx(template, output_file, dd):
    for i, gc in enumerate(dd):
        doc = DocxTemplate( template )
        doc.render(gc)
        doc.save(output_file)

        if i==0:
            master = Document(output_file)
            composer = Composer(master)
        else:
            doc = Document(output_file)
            composer.append(doc)    
    composer.save(output_file)

def abs_tmp_filename(target_filename):
    tmp_media = settings.MEDIA_TMP_EXPORT  
    abs_filename = os.path.join(tmp_media, target_filename) 
    return abs_filename

def render_html_to_pdf(template_name, context):
    html_string = render_to_string(template_name, context)
    pdf_file = io.BytesIO()
    HTML(string=html_string).write_pdf(target=pdf_file)
    pdf_file.seek(0)
    return pdf_file

def append_pdf(writer, pdf_file):
    reader = PdfReader(pdf_file)
  
    for page in reader.pages:
        writer.add_page(page)

def append_filefield(writer, filefield):
    """
    Append every page of filefield's PDF content to writer.

    Returns True if the pages were appended (or there was no file to
    append), False if the file exists but couldn't be read as a PDF - the
    caller can then log/report the problem instead of the whole export
    blowing up on one bad attachment.
    """
    if not filefield:
        return True

    try:
        f=filefield.open("rb")
    except Exception:
        logger.exception("Could not open file %r for PDF export", getattr(filefield, "name", None))
        return False

    try:
        reader = PdfReader(f)
        for page in reader.pages:
            if "/Annots" in page:
                new_annots = ArrayObject()

                for ref in page["/Annots"]:
                    annot = ref.get_object()
                    if annot.get("/Subtype") != "/Link":
                        new_annots.append(ref)

                if new_annots:
                    page[NameObject("/Annots")] = new_annots
                else:
                    del page["/Annots"]
            writer.add_page(page)
    except Exception:
        logger.exception("Could not read file %r as PDF for export", getattr(filefield, "name", None))
        return False
    finally:
        f.close()

    return True

DEFAULT_CELL_ALIGNMENT = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True
)

def xlsx_response(data, headers = None, cell_alignment = DEFAULT_CELL_ALIGNMENT, default_width = 30, widths = None, title='Sheet'):

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet") 
        )

    response["Content-Disposition"] = ( 'attachment; filename="export.xlsx"' )

    wb = Workbook()
    ws = wb.active
    ws.title = title

    if (not headers) and ( len(data) >= 1 ):
        headers = [ sanitize_excel_string(h) for h in data[0].keys() ]

    headers = headers or []

    print('Headers : ',headers)
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        
    for d in data:
        row = [sanitize_excel_string(d[k]) for k in headers]
        ws.append(row)
        
    ws.auto_filter.ref = ws.dimensions
    
    for idx, header in enumerate(headers, start=1):
        if not widths:
            ws.column_dimensions[get_column_letter(idx)].width = default_width
        else:    
            ws.column_dimensions[get_column_letter(idx)].width = widths[header]

    # Στοίχιση όλων των κελιών
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = cell_alignment

    wb.save(response)
    return response

def sanitize_excel_string(value):
    if value is None:
        return value
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))